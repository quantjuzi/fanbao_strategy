# -*- coding: utf-8 -*-
"""从东方财富电子对账单PDF中提取股票买卖明细。

只保留“证券买入”和“证券卖出”两类记录，自动忽略银证转账、
银行转证券、证券转银行、利息归本等非交易流水。
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "private"
    / "东方财富对账单_交易明细.csv"
)

TRADE_PATTERN = re.compile(
    r"^(?P<日期>\d{8})\s+"
    r"(?P<业务类型>证券买入|证券卖出)\s+"
    r"(?P<代码>\S+)\s+"
    r"(?P<名称>.+?)\s+"
    r"(?P<数量>\d+)\s+"
    r"(?P<价格>[\d.]+)\s+"
    r"(?P<发生金额>-?[\d.]+)\s+"
    r"(?P<手续费>[\d.]+)\s+"
    r"(?P<印花税>[\d.]+)\s+"
    r"(?P<过户费>[\d.]+)\s+"
    r"(?P<资金余额>-?[\d.]+)$"
)


def parse_args() -> argparse.Namespace:
    """读取命令行参数。"""

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--broker-profit",
        type=float,
        default=None,
        help="东方财富展示的净盈亏，已包含手续费。",
    )
    return parser.parse_args()


def parse_pdf(path: Path) -> pd.DataFrame:
    """提取PDF中的股票成交记录。"""

    reader = PdfReader(path)
    rows: list[dict[str, object]] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        for line in text.splitlines():
            match = TRADE_PATTERN.match(line.strip())
            if not match:
                continue
            item = match.groupdict()
            rows.append(
                {
                    "交易日期": pd.to_datetime(
                        item["日期"],
                        format="%Y%m%d",
                        errors="coerce",
                    ),
                    "方向": (
                        "买入"
                        if item["业务类型"] == "证券买入"
                        else "卖出"
                    ),
                    "证券代码": item["代码"].split(".")[0].zfill(6),
                    "证券名称": re.sub(r"\s+", "", item["名称"]),
                    "成交数量": int(item["数量"]),
                    "成交价格": round(float(item["价格"]), 4),
                    "资金发生额": round(
                        float(item["发生金额"]),
                        2,
                    ),
                    "成交金额": round(
                        abs(float(item["发生金额"])),
                        2,
                    ),
                    "手续费": round(float(item["手续费"]), 2),
                    "印花税": round(float(item["印花税"]), 2),
                    "过户费": round(float(item["过户费"]), 2),
                    "总费用": round(
                        float(item["手续费"])
                        + float(item["印花税"])
                        + float(item["过户费"]),
                        2,
                    ),
                    "资金余额": round(float(item["资金余额"]), 2),
                    "原始业务类型": item["业务类型"],
                }
            )

    result = pd.DataFrame(rows)
    if result.empty:
        raise ValueError("没有解析到股票买卖记录，请检查PDF格式。")
    result = result.drop_duplicates(
        subset=[
            "交易日期",
            "方向",
            "证券代码",
            "成交数量",
            "成交价格",
            "资金发生额",
        ],
        keep="first",
    )
    return result.sort_values(
        ["交易日期", "证券代码", "方向"]
    ).reset_index(drop=True)


def build_paired_trades(data: pd.DataFrame) -> tuple[pd.DataFrame, float]:
    """按加权平均成本配对卖出，计算已实现净盈亏。"""

    state: dict[str, dict[str, float]] = {}
    rows: list[dict[str, object]] = []
    unmatched_sell = 0.0

    ordered = data.sort_values(
        ["交易日期", "证券代码"],
        kind="stable",
    )
    for _, trade in ordered.iterrows():
        code = str(trade["证券代码"])
        shares = float(trade["成交数量"])
        price = float(trade["成交价格"])
        fee = float(trade["总费用"])
        current = state.setdefault(
            code,
            {"shares": 0.0, "cost": 0.0},
        )

        if trade["方向"] == "买入":
            current["shares"] += shares
            current["cost"] += shares * price + fee
            continue

        if current["shares"] <= 0:
            unmatched_sell += shares
            continue

        matched = min(shares, current["shares"])
        average_cost = (
            current["cost"] / current["shares"]
            if current["shares"] > 0
            else 0.0
        )
        buy_cost = average_cost * matched
        sell_fee = fee * matched / shares
        sell_amount = price * matched - sell_fee
        net_pnl = sell_amount - buy_cost
        rows.append(
            {
                "卖出日期": trade["交易日期"],
                "证券代码": code,
                "证券名称": trade["证券名称"],
                "配对数量": int(matched),
                "加权买入成本": round(average_cost, 4),
                "卖出价格": round(price, 4),
                "买入成本金额": round(buy_cost, 2),
                "卖出净金额": round(sell_amount, 2),
                "卖出费用": round(sell_fee, 2),
                "净盈亏金额": round(net_pnl, 2),
                "净收益率": (
                    round(net_pnl / buy_cost * 100, 4)
                    if buy_cost > 0
                    else 0.0
                ),
            }
        )
        current["shares"] -= matched
        current["cost"] -= buy_cost
        if shares > matched:
            unmatched_sell += shares - matched

    result = pd.DataFrame(rows)
    if not result.empty:
        result = result.sort_values(
            ["卖出日期", "证券代码"]
        ).reset_index(drop=True)
    return result, unmatched_sell


def parse_account_summary(
    path: Path,
    trades: pd.DataFrame,
) -> dict[str, float]:
    """按电子对账单资产口径估算账户总盈亏。"""

    reader = PdfReader(path)
    text = "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )
    external_pattern = re.compile(
        r"^(?P<日期>\d{8})\s+"
        r"(?P<类型>银行转证券|证券转银行)\s+"
        r"0\s+0\.0000\s+"
        r"(?P<金额>-?[\d.]+)"
    )
    external_flow = 0.0
    for line in text.splitlines():
        match = external_pattern.match(line.strip())
        if match:
            external_flow += float(match.group("金额"))

    asset_match = re.search(
        r"总资产\(RMB\)[：:]\s*([\d,.]+)",
        text,
    )
    ending_assets = (
        float(asset_match.group(1).replace(",", ""))
        if asset_match
        else float("nan")
    )

    first_trade = trades.iloc[0]
    first_balance = float(first_trade["资金余额"])
    first_cash_flow = float(first_trade["资金发生额"])
    initial_cash = first_balance - first_cash_flow
    total_profit = (
        ending_assets - initial_cash - external_flow
        if pd.notna(ending_assets)
        else float("nan")
    )
    return {
        "期初现金估算": initial_cash,
        "银证净转入": external_flow,
        "期末总资产": ending_assets,
        "账户总盈亏估算": total_profit,
        "交易现金流净额": float(trades["资金发生额"].sum()),
    }


def write_excel(
    path: Path,
    data: pd.DataFrame,
    paired: pd.DataFrame,
    account_summary: dict[str, float],
    unmatched_sell: float,
    broker_profit: float | None,
) -> None:
    """输出成交明细、汇总和月度统计。"""

    summary = pd.DataFrame(
        [
            {"指标": "成交记录", "值": len(data)},
            {"指标": "买入笔数", "值": data["方向"].eq("买入").sum()},
            {"指标": "卖出笔数", "值": data["方向"].eq("卖出").sum()},
            {
                "指标": "买入金额",
                "值": data.loc[
                    data["方向"].eq("买入"), "成交金额"
                ].sum(),
            },
            {
                "指标": "卖出金额",
                "值": data.loc[
                    data["方向"].eq("卖出"), "成交金额"
                ].sum(),
            },
            {
                "指标": "手续费合计（对账单“手续费”）",
                "值": data["手续费"].sum(),
            },
            {
                "指标": "印花税合计（对账单“印花税”）",
                "值": data["印花税"].sum(),
            },
            {
                "指标": "过户费合计（对账单“过户费”）",
                "值": data["过户费"].sum(),
            },
            {
                "指标": "总费用合计（三项相加）",
                "值": data["总费用"].sum(),
            },
        ]
    )
    if broker_profit is not None:
        summary.loc[len(summary)] = {
            "指标": "盈利（未扣费用）",
            "值": broker_profit + data["总费用"].sum(),
        }
        summary.loc[len(summary)] = {
            "指标": "扣除手续费后的总盈利",
            "值": broker_profit,
        }
    detail_view = data[
        [
            "交易日期",
            "方向",
            "证券代码",
            "证券名称",
            "成交数量",
            "成交价格",
            "成交金额",
            "总费用",
        ]
    ].copy()
    paired_view = (
        paired[
            [
                "卖出日期",
                "证券代码",
                "证券名称",
                "配对数量",
                "加权买入成本",
                "卖出价格",
                "净盈亏金额",
                "净收益率",
            ]
        ].copy()
        if not paired.empty
        else paired
    )
    fee_view = data[
        [
            "交易日期",
            "方向",
            "证券代码",
            "证券名称",
            "成交金额",
            "手续费",
            "印花税",
            "过户费",
            "总费用",
        ]
    ].copy()
    monthly = data.copy()
    monthly["月份"] = monthly["交易日期"].dt.to_period("M").astype(str)
    monthly = (
        monthly.groupby("月份")
        .agg(
            买入笔数=("方向", lambda values: values.eq("买入").sum()),
            卖出笔数=("方向", lambda values: values.eq("卖出").sum()),
            买入金额=(
                "成交金额",
                lambda values: 0.0,
            ),
            卖出金额=(
                "成交金额",
                lambda values: 0.0,
            ),
            手续费=("手续费", "sum"),
            印花税=("印花税", "sum"),
            过户费=("过户费", "sum"),
        )
        .reset_index()
    )
    for index, row in monthly.iterrows():
        month = row["月份"]
        sample = data.loc[
            data["交易日期"].dt.to_period("M").astype(str).eq(month)
        ]
        monthly.loc[index, "买入金额"] = sample.loc[
            sample["方向"].eq("买入"), "成交金额"
        ].sum()
        monthly.loc[index, "卖出金额"] = sample.loc[
            sample["方向"].eq("卖出"), "成交金额"
        ].sum()

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        detail_view.to_excel(writer, sheet_name="成交明细", index=False)
        paired_view.to_excel(writer, sheet_name="配对明细", index=False)
        fee_view.to_excel(writer, sheet_name="费用核对", index=False)
        summary.to_excel(writer, sheet_name="汇总", index=False)
        monthly.to_excel(writer, sheet_name="月度统计", index=False)

    workbook = load_workbook(path)
    paired_sheet = workbook["配对明细"]
    paired_headers = {
        cell.value: cell.column
        for cell in paired_sheet[1]
    }
    for row in paired_sheet.iter_rows(min_row=2):
        if "证券代码" in paired_headers:
            row[paired_headers["证券代码"] - 1].number_format = "@"
    detail = workbook["成交明细"]
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    for cell in detail[1]:
        cell.fill = PatternFill("solid", fgColor="24445C")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    headers = {cell.value: cell.column for cell in detail[1]}
    for row in detail.iter_rows(min_row=2):
        row[headers["交易日期"] - 1].number_format = "yyyy-mm-dd"
        row[headers["证券代码"] - 1].number_format = "@"
        row[headers["成交数量"] - 1].number_format = "#,##0"
        row[headers["成交价格"] - 1].number_format = "0.0000"
        for name in ["成交金额", "总费用"]:
            row[headers[name] - 1].number_format = "#,##0.00"
        direction = row[headers["方向"] - 1].value
        color = "C00000" if direction == "买入" else "2E7D32"
        row[headers["方向"] - 1].font = Font(
            color=color,
            bold=True,
        )

    widths = {
        "交易日期": 12,
        "方向": 8,
        "证券代码": 12,
        "证券名称": 14,
        "成交数量": 12,
        "成交价格": 12,
        "成交金额": 14,
        "总费用": 10,
    }
    for name, width in widths.items():
        detail.column_dimensions[
            chr(64 + headers[name])
        ].width = width

    fee_sheet = workbook["费用核对"]
    fee_headers = {cell.value: cell.column for cell in fee_sheet[1]}
    for row in fee_sheet.iter_rows(min_row=2):
        row[fee_headers["交易日期"] - 1].number_format = "yyyy-mm-dd"
        row[fee_headers["证券代码"] - 1].number_format = "@"
        for name in ["成交金额", "手续费", "印花税", "过户费", "总费用"]:
            row[fee_headers[name] - 1].number_format = "#,##0.00"

    summary_sheet = workbook["汇总"]
    for row in summary_sheet.iter_rows(min_row=2):
        row[1].number_format = "#,##0.00"

    monthly_sheet = workbook["月度统计"]
    for row in monthly_sheet.iter_rows(min_row=2):
        for cell in row[1:]:
            cell.number_format = "#,##0.00"

    for sheet in ["配对明细", "费用核对", "汇总", "月度统计"]:
        current = workbook[sheet]
        current.freeze_panes = "A2"
        current.auto_filter.ref = current.dimensions
        for cell in current[1]:
            cell.fill = PatternFill("solid", fgColor="24445C")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for column_cells in current.columns:
            width = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            current.column_dimensions[
                column_cells[0].column_letter
            ].width = min(max(width + 2, 10), 18)

    workbook.save(path)


def write_markdown_view(
    path: Path,
    data: pd.DataFrame,
    paired: pd.DataFrame,
    account_summary: dict[str, float],
    broker_profit: float | None,
) -> None:
    """生成适合GitHub网页阅读的Markdown表格。"""

    summary_rows = [
        ("成交记录", len(data)),
        ("买入笔数", int(data["方向"].eq("买入").sum())),
        ("卖出笔数", int(data["方向"].eq("卖出").sum())),
        ("手续费合计", data["手续费"].sum()),
        ("印花税合计", data["印花税"].sum()),
        ("过户费合计", data["过户费"].sum()),
        ("总费用合计", data["总费用"].sum()),
    ]
    if broker_profit is not None:
        summary_rows.append(
            ("盈利（未扣费用）", broker_profit + data["总费用"].sum())
        )
        summary_rows.append(
            ("扣除手续费后的总盈利", broker_profit)
        )

    preview = data[
        [
            "交易日期",
            "方向",
            "证券代码",
            "证券名称",
            "成交数量",
            "成交价格",
            "成交金额",
            "总费用",
        ]
    ].copy()

    lines = [
        "# 实盘成交明细展示",
        "",
        "本页是便于在 GitHub 直接阅读的展示版。完整数据请查看 CSV 或下载 Excel。",
        "",
        "## 汇总",
        "",
        "| 指标 | 数值 |",
        "|---|---:|",
    ]
    lines.extend(
        f"| {name} | {value:,.2f} |"
        for name, value in summary_rows
    )
    lines.extend(
        [
            "",
        "## 完整成交明细",
            "",
            "| 交易日期 | 方向 | 证券代码 | 证券名称 | 成交数量 | 成交价格 | 成交金额 | 总费用 |",
            "|---|---|---|---|---:|---:|---:|---:|",
        ]
    )
    for _, row in preview.iterrows():
        lines.append(
            "| {date} | {direction} | {code} | {name} | "
            "{quantity:,} | {price:.4f} | {amount:,.2f} | {fee:,.2f} |".format(
                date=row["交易日期"].strftime("%Y-%m-%d"),
                direction=row["方向"],
                code=row["证券代码"],
                name=row["证券名称"],
                quantity=int(row["成交数量"]),
                price=row["成交价格"],
                amount=row["成交金额"],
                fee=row["总费用"],
            )
        )
    lines.extend(
        [
            "",
            "## 文件入口",
            "",
            "- `结果/实盘成交明细_脱敏.csv`",
            "- `结果/实盘成交明细_脱敏.xlsx`",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    """解析PDF并保存脱敏成交明细。"""

    args = parse_args()
    result = parse_pdf(args.pdf)
    paired, unmatched_sell = build_paired_trades(result)
    account_summary = parse_account_summary(args.pdf, result)
    public_result = result.drop(columns=["资金余额"])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    public_result.to_csv(
        args.output,
        index=False,
        encoding="utf-8-sig",
    )
    excel_path = args.output.with_suffix(".xlsx")
    write_excel(
        excel_path,
        public_result,
        paired,
        account_summary,
        unmatched_sell,
        args.broker_profit,
    )
    markdown_path = args.output.with_suffix(".md")
    write_markdown_view(
        markdown_path,
        public_result,
        paired,
        account_summary,
        args.broker_profit,
    )

    print(
        f"解析完成：{len(result)} 条成交记录，"
        f"买入 {(result['方向'] == '买入').sum()} 条，"
        f"卖出 {(result['方向'] == '卖出').sum()} 条"
    )
    print(f"日期范围：{result['交易日期'].min()} 至 "
          f"{result['交易日期'].max()}")
    print(f"输出文件：{args.output}")
    print(f"Excel文件：{excel_path}")
    print(f"Markdown文件：{markdown_path}")
    print(
        f"已实现净盈亏：{paired['净盈亏金额'].sum():.2f}"
        if not paired.empty
        else "已实现净盈亏：0.00"
    )
    print(f"总费用：{result['总费用'].sum():.2f}")
    if args.broker_profit is not None:
        print(f"东方财富净盈亏（含费用）：{args.broker_profit:.2f}")


if __name__ == "__main__":
    main()
