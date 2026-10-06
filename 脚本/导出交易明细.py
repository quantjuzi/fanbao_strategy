# -*- coding: utf-8 -*-
'''导出已复核的67笔交易清单。'''

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE = Path(
    r'C:\Users\Administrator\PyCharmMiscProject\高开均线完整输出'
    r'\年化最高版_67笔交易清单.csv'
)
OUTPUT_DIR = Path(__file__).resolve().parents[1] / '结果'
CSV_PATH = OUTPUT_DIR / '交易明细.csv'
EXCEL_PATH = OUTPUT_DIR / '交易明细.xlsx'


def main() -> None:
    """生成CSV和美化后的Excel。"""

    data = pd.read_csv(SOURCE, encoding='utf-8-sig')
    data['买入/卖出日期'] = (
        data['模拟买入日'].str[5:]
        + '/'
        + data['模拟卖出日'].str[5:]
    )
    public = data[
        [
            '代码',
            '买入/卖出日期',
            '模拟买入价',
            '模拟卖出价',
            '净收益率',
            '净盈亏金额',
        ]
    ].copy()

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    public.to_csv(CSV_PATH, index=False, encoding='utf-8-sig')
    public.to_excel(EXCEL_PATH, index=False)

    workbook = load_workbook(EXCEL_PATH)
    worksheet = workbook.active
    worksheet.title = '交易清单'
    worksheet.freeze_panes = 'A2'
    worksheet.auto_filter.ref = worksheet.dimensions

    header_fill = PatternFill('solid', fgColor='24445C')
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center')

    for row in worksheet.iter_rows(min_row=2):
        row[0].alignment = Alignment(horizontal='center')
        row[1].alignment = Alignment(horizontal='center')
        row[2].number_format = '0.0000'
        row[3].number_format = '0.0000'
        row[4].number_format = '0.0000'
        row[5].number_format = '#,##0.00'
        if row[4].value >= 0:
            row[4].font = Font(color='C00000')
            row[5].font = Font(color='C00000')
        else:
            row[4].font = Font(color='2E7D32')
            row[5].font = Font(color='2E7D32')

    widths = [16, 18, 13, 13, 13, 15]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[
            chr(64 + index)
        ].width = width

    workbook.save(EXCEL_PATH)
    print(f'CSV：{CSV_PATH}')
    print(f'Excel：{EXCEL_PATH}')


if __name__ == '__main__':
    main()
