# -*- coding: utf-8 -*-
"""把两个入选策略整理成统一格式的逐笔交易清单。

输出包含代码、买入日期、买入价格、卖出日期、卖出价格、净收益率和净盈亏金额，
便于逐笔复核选股和成交假设。
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "结果"
SOURCE_DIR = Path(
    r"C:\Users\Administrator\PyCharmMiscProject\高开均线完整输出"
)

STRATEGY1_SOURCE = SOURCE_DIR / "年化最高版_67笔交易清单.csv"
STRATEGY2_SOURCE = SOURCE_DIR / "单板反包_开板明细.csv"

PUBLIC_COLUMNS = [
    "策略",
    "代码",
    "买入日期",
    "买入价格",
    "卖出日期",
    "卖出价格",
    "净收益率",
    "净盈亏金额",
]


def normalize_code(values: pd.Series) -> pd.Series:
    """把聚宽代码转换成6位股票代码。"""

    return (
        values.astype(str)
        .str.split(".")
        .str[0]
        .str.zfill(6)
    )


def read_strategy1() -> pd.DataFrame:
    """读取67笔连板断板策略明细。"""

    data = pd.read_csv(STRATEGY1_SOURCE, encoding="utf-8-sig")
    return pd.DataFrame(
        {
            "策略": "策略一：连板断板承接",
            "代码": normalize_code(data["代码"]),
            "买入日期": pd.to_datetime(data["模拟买入日"]),
            "买入价格": data["实际模拟买入成交价"],
            "卖出日期": pd.to_datetime(data["模拟卖出日"]),
            "卖出价格": data["实际模拟卖出成交价"],
            "净收益率": data["净收益率"],
            "净盈亏金额": data["净盈亏金额"],
        }
    )


def read_strategy2() -> pd.DataFrame:
    """读取单板反包开板2次策略明细。"""

    data = pd.read_csv(STRATEGY2_SOURCE, encoding="utf-8-sig")
    data = data.loc[data["开板次数"].eq(2)].copy()
    return pd.DataFrame(
        {
            "策略": "策略二：单板反包开板2次",
            "代码": normalize_code(data["code"]),
            "买入日期": pd.to_datetime(data["买入日期"]),
            "买入价格": data["买入成交价"],
            "卖出日期": pd.to_datetime(data["卖出日期"]),
            "卖出价格": data["卖出成交价"],
            "净收益率": data["净收益率"],
            "净盈亏金额": data["净盈亏金额"],
        }
    )


def summarize(
    name: str,
    trades: pd.DataFrame,
) -> dict[str, float | str]:
    """汇总单个策略的绩效指标。"""

    winners = trades.loc[trades["净收益率"].gt(0), "净收益率"]
    losers = trades.loc[trades["净收益率"].lt(0), "净收益率"]
    return {
        "策略": name,
        "交易数": len(trades),
        "净胜率": trades["净收益率"].gt(0).mean() * 100,
        "平均净收益率": trades["净收益率"].mean(),
        "平均盈利": winners.mean() if not winners.empty else 0.0,
        "平均亏损": losers.mean() if not losers.empty else 0.0,
        "盈亏比": (
            winners.mean() / abs(losers.mean())
            if not winners.empty and not losers.empty
            else 0.0
        ),
        "按两万元每笔累计盈亏": (
            trades["净收益率"].mean()
            * len(trades)
            * 20_000
            / 100
        ),
        "实际整手累计盈亏": trades["净盈亏金额"].sum(),
    }


def concentration_summary(
    name: str,
    trades: pd.DataFrame,
) -> dict[str, float | str]:
    """统计收益是否集中在少数大额盈利交易上。"""

    sample = trades.copy()
    sample["标准化盈亏"] = (
        sample["净收益率"] * 20_000 / 100
    )
    total_pnl = float(sample["标准化盈亏"].sum())
    sorted_pnl = sample["标准化盈亏"].sort_values(
        ascending=False
    )

    result: dict[str, float | str] = {
        "策略": name,
        "交易数": len(sample),
        "累计标准化盈亏": total_pnl,
        "中位数收益率": sample["净收益率"].median(),
        "最大单笔盈利": sample["净收益率"].max(),
        "最大单笔亏损": sample["净收益率"].min(),
    }
    for count in [1, 3, 5]:
        removed = sorted_pnl.head(count)
        remaining = sample.drop(removed.index)
        result[f"前{count}笔盈利贡献"] = (
            removed.sum() / total_pnl * 100
            if total_pnl != 0
            else 0.0
        )
        result[f"去掉前{count}笔后累计盈亏"] = remaining[
            "标准化盈亏"
        ].sum()
        result[f"去掉前{count}笔后平均收益率"] = (
            remaining["标准化盈亏"].sum()
            / len(remaining)
            / 20_000
            * 100
            if not remaining.empty
            else 0.0
        )
    return result


def write_excel(path: Path, trades: pd.DataFrame) -> None:
    """输出带格式的Excel明细。"""

    output = trades.copy()
    output["买入日期"] = output["买入日期"].dt.strftime("%Y-%m-%d")
    output["卖出日期"] = output["卖出日期"].dt.strftime("%Y-%m-%d")
    output.to_excel(path, index=False)

    workbook = load_workbook(path)
    worksheet = workbook.active
    worksheet.title = "交易明细"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="24445C")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }
    for row in worksheet.iter_rows(min_row=2):
        row[headers["代码"] - 1].number_format = "@"
        row[headers["买入价格"] - 1].number_format = "0.0000"
        row[headers["卖出价格"] - 1].number_format = "0.0000"
        row[headers["净收益率"] - 1].number_format = "0.0000"
        row[headers["净盈亏金额"] - 1].number_format = "#,##0.00"
        color = (
            "C00000"
            if row[headers["净收益率"] - 1].value >= 0
            else "2E7D32"
        )
        row[headers["净收益率"] - 1].font = Font(color=color)
        row[headers["净盈亏金额"] - 1].font = Font(color=color)

    workbook.save(path)


def main() -> None:
    """导出两个策略的CSV、Excel和绩效汇总。"""

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    strategy1 = read_strategy1()
    strategy2 = read_strategy2()

    strategy1_path = OUTPUT_DIR / "策略一_连板断板_交易明细"
    strategy2_path = OUTPUT_DIR / "策略二_单板反包_交易明细"

    strategy1.to_csv(
        strategy1_path.with_suffix(".csv"),
        index=False,
        encoding="utf-8-sig",
    )
    strategy2.to_csv(
        strategy2_path.with_suffix(".csv"),
        index=False,
        encoding="utf-8-sig",
    )
    write_excel(strategy1_path.with_suffix(".xlsx"), strategy1)
    write_excel(strategy2_path.with_suffix(".xlsx"), strategy2)

    summary = pd.DataFrame(
        [
            summarize("策略一：连板断板承接", strategy1),
            summarize("策略二：单板反包开板2次", strategy2),
        ]
    )
    summary.to_csv(
        OUTPUT_DIR / "两个策略绩效汇总.csv",
        index=False,
        encoding="utf-8-sig",
    )
    concentration = pd.DataFrame(
        [
            concentration_summary(
                "策略一：连板断板承接",
                strategy1,
            ),
            concentration_summary(
                "策略二：单板反包开板2次",
                strategy2,
            ),
        ]
    )
    concentration.to_csv(
        OUTPUT_DIR / "两个策略盈利集中度.csv",
        index=False,
        encoding="utf-8-sig",
    )
    print(summary.to_string(index=False))
    print()
    print(concentration.to_string(index=False))


if __name__ == "__main__":
    main()
